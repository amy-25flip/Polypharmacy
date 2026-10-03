"""Turn a returned clinician review workbook into doctor_review.json.

Usage: python scripts/disease_formulary/apply_doctor_review.py path/to/returned.xlsx
Then rebuild the lists: python scripts/disease_formulary/build_disease_formulary.py

Only explicit decisions change anything: "Remove" drops a medicine from a disease's list and
rows on "Add medicines" add one (credited to the reviewer). Keep, Unsure and blank rows leave
the list as it is. Requests for diseases or medicines PolyGuard does not know are reported,
never applied.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).with_name("doctor_review.json")


def cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def main(path: Path) -> None:
    wb = load_workbook(path, data_only=True)
    vocabulary = {n.casefold(): n for n in json.loads((ROOT / "processed" / "drug_vocabulary.json").read_text(encoding="utf-8"))}
    diseases = {d["name"].casefold(): d["name"]
                for d in json.loads((ROOT / "processed" / "disease_formulary.json").read_text(encoding="utf-8"))["diseases"]}
    reviewer = cell_text(wb["Read me"]["B4"].value) or "unnamed reviewer"
    review_date = cell_text(wb["Read me"]["B5"].value)

    removed, unsure, kept, requests = [], [], 0, []
    for row in wb["Review"].iter_rows(min_row=2, values_only=True, max_col=9):
        disease, medicine, decision, note = cell_text(row[0]), cell_text(row[1]), cell_text(row[7]), cell_text(row[8])
        if not disease or not medicine:
            continue
        if decision == "Remove":
            removed.append({"disease": disease, "medicine": medicine, "note": note})
        elif decision == "Unsure":
            unsure.append({"disease": disease, "medicine": medicine, "note": note})
        elif decision == "Keep":
            kept += 1

    added = []
    for row in wb["Add medicines"].iter_rows(min_row=2, values_only=True, max_col=4):
        disease, medicine, reason, note = (cell_text(v) for v in row)
        if not disease and not medicine:
            continue
        if disease.casefold() not in diseases:
            requests.append(f"New disease requested: '{disease}' (medicine '{medicine}'). {reason}".strip())
        elif medicine.casefold() not in vocabulary:
            requests.append(f"Medicine not in PolyGuard's drug list: '{medicine}' for {diseases[disease.casefold()]}. {reason}".strip())
        else:
            added.append({"disease": diseases[disease.casefold()], "medicine": vocabulary[medicine.casefold()],
                          "reason": reason, "note": note})

    OUTPUT.write_text(json.dumps({"reviewer": reviewer, "review_date": review_date, "removed": removed,
                                  "added": added, "unsure": unsure, "requests_not_applied": requests},
                                 ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Reviewer: {reviewer} {review_date}".strip())
    print(f"Keep: {kept}  Remove: {len(removed)}  Unsure: {len(unsure)}  Added: {len(added)}  "
          f"Requests not applied: {len(requests)}")
    for item in requests:
        print(f"  - {item}")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]))
