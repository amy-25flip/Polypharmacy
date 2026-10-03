"""Build the clinician review workbook for the disease -> medicine reference lists.

Run from any directory: python scripts/disease_formulary/make_review_workbook.py
Output: doctor_review/PolyGuard_Disease_Medicine_Review.xlsx

The doctor marks each medicine Keep / Remove / Unsure, lists missing medicines on the
"Add medicines" sheet, and sends the file back. apply_doctor_review.py turns it into
doctor_review.json, which build_disease_formulary.py applies on the next rebuild.
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
FORMULARY = ROOT / "processed" / "disease_formulary.json"
CURATED = Path(__file__).with_name("curated_formulary.json")
VOCABULARY = ROOT / "processed" / "drug_vocabulary.json"
OUTPUT = ROOT / "doctor_review" / "PolyGuard_Disease_Medicine_Review.xlsx"

ADD_ROWS = 300
HETIONET_LABELS = {"hetionet:CtD": "Treats", "hetionet:CpD": "Symptom relief"}

NAVY = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def set_widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def build() -> None:
    formulary = json.loads(FORMULARY.read_text(encoding="utf-8"))["diseases"]
    curated = json.loads(CURATED.read_text(encoding="utf-8"))
    urls = curated["sources"]
    vocabulary = sorted(json.loads(VOCABULARY.read_text(encoding="utf-8")), key=str.casefold)
    spread = Counter(m["name"] for d in formulary for m in d["medicines"])

    rows = []
    for disease in formulary:
        for med in disease["medicines"]:
            guideline = [s for s in med["sources"] if not s.startswith("hetionet:")]
            database = [HETIONET_LABELS.get(s, s) for s in med["sources"] if s.startswith("hetionet:")]
            if guideline:
                level = "Guideline-cited" + (" + database" if database else "")
            else:
                level = "Database only"
            flags = []
            if not guideline:
                flags.append("No guideline citation: please confirm")
            if spread[med["name"]] >= 10:
                flags.append(f"Listed under {spread[med['name']]} diseases: may be a general-purpose drug")
            rows.append([disease["name"], med["name"], level,
                         "\n".join(guideline) or "-",
                         "\n".join(dict.fromkeys(urls[s] for s in guideline if s in urls)) or "-",
                         ", ".join(database) or "-", "; ".join(flags) or "-"])

    wb = Workbook()

    # ---------------- Read me ----------------
    readme = wb.active
    readme.title = "Read me"
    readme.sheet_view.showGridLines = False
    set_widths(readme, [26, 90])
    readme["A1"] = "PolyGuard: disease and medicine list review"
    readme["A1"].font = Font(bold=True, size=16, color=NAVY)
    readme["A2"] = ("These lists let a doctor pick a diagnosis and browse medicines associated with it. "
                    "They support browsing only. They are not prescribing advice. Your review decides "
                    "which medicines appear.")
    readme["A2"].alignment = WRAP
    readme.merge_cells("A2:B2")
    readme.row_dimensions[2].height = 34
    readme["A4"], readme["A5"] = "Reviewer name", "Review date"
    readme["B4"].fill = readme["B5"].fill = INPUT_FILL
    for ref in ("A4", "A5"):
        readme[ref].font = Font(bold=True)
    for ref in ("B4", "B5"):
        readme[ref].border = BORDER
    steps = [
        ("How to review", ""),
        ("1. Open the Review sheet", "One row per disease and medicine pair. Use the filter arrows to look at one disease, or at "
                                     "only the rows marked \"Database only\"."),
        ("2. Fill the Decision column", "Choose Keep (appropriate to list for this disease), Remove (should not be listed), or "
                                        "Unsure (leave it listed, flag it for a colleague). A blank row counts as not yet reviewed "
                                        "and stays listed."),
        ("3. Add what is missing", "On the Add medicines sheet, choose the disease and medicine from the drop-downs and give a short "
                                   "reason. If a disease or medicine is not in the drop-down, type it in anyway. It will be reported "
                                   "back as a request."),
        ("4. Track progress", "The Diseases sheet counts how many medicines you have reviewed for each disease. "
                              "Set its Status column to Reviewed once a disease is finished."),
        ("5. Send the file back", "Fill in the reviewer name and date above and return the workbook."),
        ("", ""),
        ("Evidence level", ""),
        ("Guideline-cited", "Listed in a named guideline or essential-medicines list (NLEM 2022, WHO EML, NICE, ESC, KDIGO, "
                            "GINA, GOLD, India MoHFW). The citation and link are on the row."),
        ("Database only", "Appears only in the public Hetionet knowledge graph (\"Treats\" or \"Symptom relief\" links). There is no "
                          "guideline behind it, so it needs the closest look. Some of these are not first-line or are off-label."),
        ("Broad drug flag", "A medicine listed under 10 or more diseases is often a general-purpose drug "
                            "(a corticosteroid, for example). Check that it belongs on that disease's list."),
        ("", ""),
        ("What a decision changes", "Only which medicines are offered under each disease in the Plan by diagnosis tab. "
                                    "Interaction checking works on any medicine and is unaffected."),
    ]
    row = 7
    for title, text in steps:
        readme.cell(row=row, column=1, value=title).font = Font(bold=True, color=NAVY if not text else "000000",
                                                                size=13 if not text and title else 11)
        cell = readme.cell(row=row, column=2, value=text)
        cell.alignment = WRAP
        readme.cell(row=row, column=1).alignment = WRAP
        if text:
            readme.row_dimensions[row].height = max(18, 16 * (len(text) // 85 + 1))
        row += 1

    # ---------------- Diseases ----------------
    ds = wb.create_sheet("Diseases")
    set_widths(ds, [44, 12, 16, 16, 16, 16, 14, 50])
    heads = ["Disease", "Medicines", "Guideline-cited", "Database only", "Reviewed", "Marked Remove", "Status", "Reviewer notes"]
    for col, head in enumerate(heads, start=1):
        ds.cell(row=1, column=col, value=head)
    style_header(ds, 1, len(heads))
    last = len(rows) + 1
    for index, disease in enumerate(formulary, start=2):
        name = disease["name"]
        ds.cell(row=index, column=1, value=name)
        ds.cell(row=index, column=2, value=f"=COUNTIF(Review!$A$2:$A${last},A{index})")
        ds.cell(row=index, column=3, value=f'=COUNTIFS(Review!$A$2:$A${last},A{index},Review!$C$2:$C${last},"Guideline-cited*")')
        ds.cell(row=index, column=4, value=f'=COUNTIFS(Review!$A$2:$A${last},A{index},Review!$C$2:$C${last},"Database only")')
        ds.cell(row=index, column=5, value=f'=COUNTIFS(Review!$A$2:$A${last},A{index},Review!$H$2:$H${last},"?*")')
        ds.cell(row=index, column=6, value=f'=COUNTIFS(Review!$A$2:$A${last},A{index},Review!$H$2:$H${last},"Remove")')
        for col in range(1, 9):
            cell = ds.cell(row=index, column=col)
            cell.border, cell.font = BORDER, BODY_FONT
        for col in (7, 8):
            ds.cell(row=index, column=col).fill = INPUT_FILL
    disease_last = len(formulary) + 1
    status_dv = DataValidation(type="list", formula1='"Not started,In progress,Reviewed"', allow_blank=True)
    ds.add_data_validation(status_dv)
    status_dv.add(f"G2:G{disease_last}")
    ds.conditional_formatting.add(f"G2:G{disease_last}", CellIsRule(operator="equal", formula=['"Reviewed"'],
                                                                     fill=PatternFill("solid", bgColor="C6EFCE")))
    ds.conditional_formatting.add(f"G2:G{disease_last}", CellIsRule(operator="equal", formula=['"In progress"'],
                                                                     fill=PatternFill("solid", bgColor="FFEB9C")))
    ds.freeze_panes = "B2"
    ds.auto_filter.ref = f"A1:H{disease_last}"

    # ---------------- Review ----------------
    rv = wb.create_sheet("Review")
    heads = ["Disease", "Medicine", "Evidence level", "Guideline source", "Source link", "Database link type",
             "Flag", "Decision", "Doctor's notes"]
    for col, head in enumerate(heads, start=1):
        rv.cell(row=1, column=col, value=head)
    style_header(rv, 1, len(heads))
    rv.cell(row=1, column=8).fill = rv.cell(row=1, column=9).fill = PatternFill("solid", fgColor="B45309")
    for index, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            cell = rv.cell(row=index, column=col, value=value)
            cell.alignment, cell.border, cell.font = WRAP, BORDER, BODY_FONT
        for col in (8, 9):
            cell = rv.cell(row=index, column=col)
            cell.fill, cell.border, cell.alignment = INPUT_FILL, BORDER, WRAP
    set_widths(rv, [34, 24, 20, 42, 46, 18, 40, 14, 42])
    rv.freeze_panes = "C2"
    rv.auto_filter.ref = f"A1:I{last}"
    decision_dv = DataValidation(type="list", formula1='"Keep,Remove,Unsure"', allow_blank=True,
                                 showErrorMessage=True, errorTitle="Decision",
                                 error="Choose Keep, Remove or Unsure.")
    rv.add_data_validation(decision_dv)
    decision_dv.add(f"H2:H{last}")
    for value, colour in (("Keep", "C6EFCE"), ("Remove", "FFC7CE"), ("Unsure", "FFEB9C")):
        rv.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=[f'"{value}"'],
                                                                fill=PatternFill("solid", bgColor=colour)))
    rv.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=['"Database only"'],
                                                            fill=PatternFill("solid", bgColor="FDE9D9")))

    # ---------------- Add medicines ----------------
    ad = wb.create_sheet("Add medicines")
    heads = ["Disease", "Medicine to add", "Reason or source (guideline, textbook, experience)", "Doctor's notes"]
    for col, head in enumerate(heads, start=1):
        ad.cell(row=1, column=col, value=head)
    style_header(ad, 1, len(heads))
    set_widths(ad, [40, 30, 56, 40])
    for index in range(2, ADD_ROWS + 2):
        for col in range(1, 5):
            cell = ad.cell(row=index, column=col)
            cell.fill, cell.border, cell.alignment = INPUT_FILL, BORDER, WRAP
    # Warning-style validation: the lists guide the choice but never block free text,
    # so a missing disease or drug can still be typed and reported back.
    disease_dv = DataValidation(type="list", formula1=f"=Diseases!$A$2:$A${disease_last}", allow_blank=True,
                                errorStyle="warning", showErrorMessage=True, errorTitle="Not in the list",
                                error="This disease is not in the current list. Keep it anyway? It will be passed back as a request for a new disease.")
    drug_dv = DataValidation(type="list", formula1=f"='Drug list'!$A$2:$A${len(vocabulary) + 1}", allow_blank=True,
                             errorStyle="warning", showErrorMessage=True, errorTitle="Not in the list",
                             error="PolyGuard does not know this medicine yet. Keep it anyway? It will be passed back as a request.")
    ad.add_data_validation(disease_dv)
    ad.add_data_validation(drug_dv)
    disease_dv.add(f"A2:A{ADD_ROWS + 1}")
    drug_dv.add(f"B2:B{ADD_ROWS + 1}")
    ad.freeze_panes = "A2"

    # ---------------- Not available ----------------
    skipped = wb.create_sheet("Not yet available")
    set_widths(skipped, [110])
    skipped["A1"] = "Medicines and diseases left out because PolyGuard cannot check them yet"
    skipped["A1"].fill, skipped["A1"].font = HEADER_FILL, HEADER_FONT
    skipped["A2"] = "For information. If one matters for your patients, say so in the notes and it can be added to the checker's drug list."
    skipped["A2"].alignment = WRAP
    for index, item in enumerate(curated.get("skipped_for_review", []), start=3):
        skipped.cell(row=index, column=1, value=item).alignment = WRAP

    # ---------------- Drug list (validation source) ----------------
    dl = wb.create_sheet("Drug list")
    dl["A1"] = "Medicines PolyGuard can check"
    dl["A1"].fill, dl["A1"].font = HEADER_FILL, HEADER_FONT
    for index, name in enumerate(vocabulary, start=2):
        dl.cell(row=index, column=1, value=name)
    dl.column_dimensions["A"].width = 40
    dl.freeze_panes = "A2"

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Review", "Add medicines") else "9AA5B1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    db_only = sum(1 for r in rows if r[2] == "Database only")
    print(f"PASS: {len(formulary)} diseases, {len(rows)} medicine rows ({db_only} database-only) -> {OUTPUT}")


if __name__ == "__main__":
    build()
